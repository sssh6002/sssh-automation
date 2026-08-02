# -*- coding: utf-8 -*-
"""
review_sheet.py
公文批次審核表（桌面「公告彙整.xlsx」）的讀寫。

流程上的定位 — 兩道關卡都由使用者說了算:

    備料(main.py 4) → 本模組填「文號/主旨/摘要/擬辦」
                    → 你審,在「陳會」欄打 OK
                    → 批次貼擬辦送陳核(3/4 批次實作)

    公告文案產生    → 本模組填「公告」欄
                    → 你審,在「張貼」欄打 OK
                    → 批次貼校網(4/4 批次實作)

欄位固定為 COLUMNS,順序與使用者提供的範例檔一致。

重要原則:**永遠不覆蓋非空白儲存格**。使用者手改過的內容(潤過的摘要、自己重寫
的擬辦)一律保留,程式只補空白。要強制覆寫得明確傳 overwrite=True。

摘要欄不叫 LLM — 實測使用者範例的「摘要」與 內容.txt 的「說明」段落逐字相同,
直接抄即可,省一趟 API。
"""

import glob
import os
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

_BASE_DIR = os.path.dirname(os.path.abspath(__file__))

COLUMNS = ["文號", "主旨", "摘要", "擬辦", "陳會", "公告", "張貼"]
DOC_NO_COL = "文號"
GATES = ("陳會", "張貼")          # 兩道人工關卡欄位
APPROVE_TOKENS = {"ok", "o", "y", "yes", "v", "是", "可"}
# 「這筆我自己已經辦掉了,程式不要碰」。與 APPROVE 明確區分:
#   空白   = 還沒審,等你決定
#   OK     = 請程式去做
#   已辦   = 我手動做完了,不要做也不要再產文案
DONE_TOKENS = {"已辦", "已送", "已公告", "已陳核", "已存查", "手動", "跳過",
               "略過", "不用", "done", "skip", "x", "-", "—"}
DONE_MARK = "已辦"

_DOC_NO_RE = re.compile(r"(MW[A-Z]*\d+)")
_SUBJECT_RE = re.compile(r"^主旨[:：]\s*(.*)$")
_HANDLING_RE = re.compile(r"^##(?!#)\s*(.*)$")     # 承辦文字(## 但不是 ###/####)

# 掃描來源 — 與 post_web_review 同兩個工作區目錄
SCAN_DIRS = [os.path.join(_BASE_DIR, "document_download"),
             os.path.join(_BASE_DIR, "document_download_closure")]


def sheet_path():
    """審核表路徑。env.env 的 review_sheet_path 優先,否則放桌面「公告彙整.xlsx」。"""
    try:
        from taipeion_login_selenium import _read_config
        p = _read_config("review_sheet_path")
    except Exception:
        p = None
    if p:
        return os.path.abspath(os.path.expanduser(p))
    return os.path.join(os.path.expanduser("~"), "Desktop", "公告彙整.xlsx")


# ── 從公文目錄抽欄位 ───────────────────────────────────────────────────────

def _read_first(pattern_dir, pattern):
    hits = sorted(glob.glob(os.path.join(pattern_dir, pattern)))
    if not hits:
        return None
    try:
        return open(hits[0], encoding="utf-8").read()
    except OSError:
        return None


def _extract_explanation(content_txt):
    """抽 內容.txt 的「說明：」段落(到「附件：」或檔尾為止),原文不改。

    使用者範例的「摘要」欄就是這一段逐字複製 — 不要動它的用字、標點、條號。
    """
    if not content_txt:
        return None
    lines = content_txt.splitlines()
    out = []
    started = False
    for ln in lines:
        s = ln.strip()
        if not started:
            if s.startswith("說明"):
                started = True
                rest = re.sub(r"^說明[:：]\s*", "", s)
                if rest:
                    out.append(rest)
            continue
        if re.match(r"^(附件|正本|副本)[:：]", s):
            break
        if s:
            out.append(s)
    return "\n".join(out) or None


def collect(doc_dir):
    """從公文目錄抽出審核表要的四欄。回 dict;抽不到文號則回 None。

    文號    ← 目錄名裡的 MWAA 編號
    主旨    ← 內容.txt(沒有則總結.md)的「主旨：」
    摘要    ← 內容.txt 的「說明：」段落原文
    擬辦    ← 總結.md 的 ## 承辦文字
    """
    name = os.path.basename(os.path.normpath(doc_dir))
    m = _DOC_NO_RE.search(name)
    if not m:
        return None

    content = _read_first(doc_dir, "*內容.txt")
    summary = _read_first(doc_dir, "*總結.*.md")

    subject = handling = None
    for text in (content, summary):
        if not text:
            continue
        for ln in text.splitlines():
            s = ln.strip()
            if subject is None:
                ms = _SUBJECT_RE.match(s)
                if ms:
                    subject = ms.group(1).strip()
            if handling is None:
                mh = _HANDLING_RE.match(s)
                if mh and mh.group(1).strip():
                    handling = mh.group(1).strip()

    rec = {"文號": m.group(1),
           "主旨": subject,
           "摘要": _extract_explanation(content),
           "擬辦": handling,
           "_dir": os.path.abspath(doc_dir)}
    rec.update(_detect_done(doc_dir))
    return rec


def _detect_done(doc_dir):
    """從工作區的既有痕跡推出「這關已經辦過了」,免得使用者一筆一筆手標。

    - 目錄內有 *已公告.txt          → 張貼已辦
    - 目錄內有 *已陳核.txt / *已存查.txt → 陳會已辦
    - 目錄位於 document_download_closure/ → 陳會已辦
      (公文能進到待結案,代表陳核一定跑過了)

    只回「已辦」,不回空值 —— 配合 upsert 不覆蓋非空欄的規則,使用者自己打的
    OK 不會被蓋掉。
    """
    out = {}
    try:
        names = os.listdir(doc_dir)
    except OSError:
        names = []
    if any(n.endswith("已公告.txt") for n in names):
        out["張貼"] = DONE_MARK
    if any(n.endswith("已陳核.txt") or n.endswith("已存查.txt") for n in names):
        out["陳會"] = DONE_MARK
    parent = os.path.basename(os.path.dirname(os.path.normpath(doc_dir)))
    if parent.endswith("_closure"):
        out["陳會"] = DONE_MARK
    return out


def iter_doc_dirs(scan_dirs=None):
    """回所有工作區公文目錄(第一層、名稱含 MW 編號)。"""
    out = []
    for base in (scan_dirs or SCAN_DIRS):
        if not os.path.isdir(base):
            continue
        for n in sorted(os.listdir(base)):
            d = os.path.join(base, n)
            if os.path.isdir(d) and _DOC_NO_RE.search(n):
                out.append(d)
    return out


# ── xlsx 讀寫 ──────────────────────────────────────────────────────────────

def _open(path=None, create=True):
    """回 (workbook, worksheet, path)。檔案不存在且 create 則新建含表頭的檔。"""
    from openpyxl import Workbook, load_workbook
    path = path or sheet_path()
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.worksheets[0]
        if [c.value for c in ws[1]][:len(COLUMNS)] != COLUMNS:
            raise ValueError(f"{path} 第一列不是預期表頭 {COLUMNS}")
        return wb, ws, path
    if not create:
        raise FileNotFoundError(path)
    wb = Workbook()
    ws = wb.active
    ws.append(COLUMNS)
    return wb, ws, path


def _col_index(name):
    return COLUMNS.index(name) + 1


def _row_map(ws):
    """回 {文號: 列號}。"""
    out = {}
    ci = _col_index(DOC_NO_COL)
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=ci).value
        if v:
            out[str(v).strip()] = r
    return out


def _link_dir(ws, r, path):
    """把文號儲存格做成「開啟該公文資料夾」的超連結。path 不存在就不加。"""
    if not path or not os.path.isdir(path):
        return
    from openpyxl.styles import Font
    cell = ws.cell(row=r, column=_col_index(DOC_NO_COL))
    cell.hyperlink = "file:///" + os.path.abspath(path).replace("\\", "/")
    cell.font = Font(color="0563C1", underline="single")


def _style_row(ws, r):
    """新列統一設成上對齊 + 自動換行,不然長摘要會擠成一條線看不了。"""
    from openpyxl.styles import Alignment
    al = Alignment(wrap_text=True, vertical="top")
    for c in range(1, len(COLUMNS) + 1):
        ws.cell(row=r, column=c).alignment = al


def upsert(records, path=None, overwrite=False):
    """把 records(dict 或 dict list)寫進審核表。回 (新增數, 更新數)。

    - 以「文號」為 key;已存在就更新,不存在就 append。
    - 預設只填**空白**儲存格 — 你手改過的內容不會被蓋掉。overwrite=True 才強制覆寫。
    - 值為 None 的欄一律跳過。
    """
    if isinstance(records, dict):
        records = [records]
    records = [r for r in records if r and r.get(DOC_NO_COL)]
    if not records:
        return 0, 0

    wb, ws, path = _open(path)
    rows = _row_map(ws)
    added = updated = 0

    for rec in records:
        doc_no = str(rec[DOC_NO_COL]).strip()
        r = rows.get(doc_no)
        is_new = r is None
        if is_new:
            r = max(ws.max_row + 1, 2)
            ws.cell(row=r, column=_col_index(DOC_NO_COL), value=doc_no)
            _style_row(ws, r)
            rows[doc_no] = r
            added += 1
        # 文號做成資料夾超連結 —— 點一下直接開該公文目錄挑附件。
        # 每次都重設(資料夾可能被改名/搬走),不算「使用者編輯」故不受 overwrite 限制。
        _link_dir(ws, r, rec.get("_dir"))

        touched = False
        for col, val in rec.items():
            if col == DOC_NO_COL or col.startswith("_") or col not in COLUMNS \
                    or val is None:
                continue
            cell = ws.cell(row=r, column=_col_index(col))
            if cell.value not in (None, "") and not overwrite:
                continue
            cell.value = val
            touched = True
        if touched and not is_new:
            updated += 1

    _autosize(ws)
    _save(wb, path)
    return added, updated


def _save(wb, path):
    """存檔。Excel 開著時檔案被鎖 → 給人話,不要吐 traceback。"""
    try:
        wb.save(path)
    except PermissionError:
        raise PermissionError(
            f"寫不進 {path} — 這個檔正被 Excel 開著。\n"
            f"        請先關閉 Excel（或關掉該活頁簿）再重跑一次。") from None


def _autosize(ws):
    """欄寬:文號/關卡欄窄,文字欄給固定寬度(內容太長不靠欄寬,靠自動換行)。"""
    widths = {"文號": 18, "主旨": 46, "摘要": 60, "擬辦": 28,
              "陳會": 8, "公告": 60, "張貼": 8}
    from openpyxl.utils import get_column_letter
    for i, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 20)


def rows(path=None):
    """回 [{欄名: 值, '_row': 列號}, ...]。"""
    wb, ws, path = _open(path, create=False)
    out = []
    for r in range(2, ws.max_row + 1):
        rec = {name: ws.cell(row=r, column=i).value
               for i, name in enumerate(COLUMNS, start=1)}
        if not rec.get(DOC_NO_COL):
            continue
        rec["_row"] = r
        out.append(rec)
    return out


def _norm(value):
    if value is None:
        return ""
    return str(value).replace("　", " ").strip().lower()


def is_approved(value):
    """關卡欄是否算「放行」。容忍大小寫、全形空白、OK/O/Y/是 等寫法。"""
    return _norm(value) in APPROVE_TOKENS


def is_done(value):
    """關卡欄是否算「我自己已經辦掉了」。程式看到就完全不碰這一關。"""
    return _norm(value) in DONE_TOKENS


def is_archived(row, doc_dir):
    """這筆公文是否算「整個辦完了」,可以移出待辦清單、進「舊文」。

    規則(2026-07-31 承辦人決定):

        已存查　＋　（本來就不用公告　或　校網已張貼）

    **為什麼不是「兩道關卡都辦完」**:不用公告的公文「張貼」欄會永遠空白,
    拿那個當條件會讓大多數公文永遠卡在待辦清單裡。兩種公文各有各的終點 —
    不用公告的止於存查,要公告的止於張貼。

    各項判斷的來源:
    - 已存查:目錄內有 *已存查.txt(結案存查歸檔成功才寫)。
      **不能拿「陳會=已辦」代替** — 那一欄把已陳核與已存查混在一起
      (見 _detect_done),用它會讓只陳核完、還沒存查的公文提早消失。
    - 要不要公告:沿用校網張貼那支的 _should_post,不另立標準。與
      announce_doc 同一個判斷來源,否則會出現「舊文認定不用公告、公告
      批次卻還在幫它產文案」的矛盾。
    - 已張貼:**以磁碟的 *已公告.txt 為準**,審核表「張貼」欄標「已辦」也算。
      不能只看那一欄:upsert 只填空白格,使用者打過「OK」(請程式去貼)的格子
      在貼完之後不會被改寫成「已辦」,永遠停在「OK」。2026-07-31 實測
      MWAA1156007365 就是這個狀態 —— 只看欄位會讓「正常流程貼出去的公文」
      全部卡在待辦清單,等於這個功能失效。

    doc_dir 為 None / 資料夾不在了 → 無從判斷,保守回 False(留在待辦清單,
    寧可多看一眼也不要默默消失)。
    """
    if not doc_dir or not os.path.isdir(doc_dir):
        return False
    try:
        names = os.listdir(doc_dir)
    except OSError:
        return False
    if not any(n.endswith("已存查.txt") for n in names):
        return False
    if any(n.endswith("已公告.txt") for n in names) or is_done(row.get("張貼")):
        return True
    # 走到這:已存查、但還沒張貼 → 只有「本來就不用公告」才算辦完。
    try:
        from document_closure.document_closure_post_web import (
            _parse_summary, _should_post)
        return not _should_post(_parse_summary(doc_dir))
    except Exception:
        return False


def approved(gate, path=None):
    """回該關卡欄已放行的列。gate 須為 '陳會' 或 '張貼'。"""
    if gate not in GATES:
        raise ValueError(f"gate 必須是 {GATES} 之一,收到 {gate!r}")
    return [r for r in rows(path) if is_approved(r.get(gate))]


def prepare(scan_dirs=None):
    """對「還沒備料」的公文目錄補跑 LLM 摘要。回 (成功數, 失敗清單)。

    只處理已下載但缺 內容.txt／總結.md 的目錄 —— 下載本身要讀卡機,不在這裡做。
    這是 `--prepare` 的本體:讓「新公文進不來」有一個指令可以解。
    """
    from summarize_doc import summarize_doc
    ok, failed = 0, []
    todo = [d for d in iter_doc_dirs(scan_dirs)
            if not glob.glob(os.path.join(d, "*總結.*.md"))]
    if not todo:
        print("[review_sheet] 沒有待備料的公文目錄。")
        return 0, []
    print(f"[review_sheet] {len(todo)} 個目錄待備料（要叫 LLM，會花 token）")
    for i, d in enumerate(todo, 1):
        print(f"\n[{i}/{len(todo)}] {os.path.basename(d)}")
        try:
            if summarize_doc(d):
                ok += 1
            else:
                failed.append(os.path.basename(d))
        except Exception as e:
            print(f"      [ERROR] {type(e).__name__}: {e}")
            failed.append(os.path.basename(d))
    return ok, failed


def sync(scan_dirs=None, path=None, overwrite=False):
    """掃工作區公文目錄 → 寫進審核表。回 (新增, 更新, 掃到幾個, 略過幾個)。

    只寫**備料完成**(抓得到主旨)的目錄。只有來文、還沒跑摘要的空殼一律略過 —
    生一列只有文號的空白列對審核沒有幫助,只是讓表變髒。
    """
    dirs = iter_doc_dirs(scan_dirs)
    recs, skipped = [], 0
    for d in dirs:
        c = collect(d)
        if not c:
            skipped += 1
        elif not c.get("主旨"):
            skipped += 1
        else:
            recs.append(c)
    added, updated = upsert(recs, path=path, overwrite=overwrite)
    return added, updated, len(dirs), skipped


def main():
    import argparse
    ap = argparse.ArgumentParser(description="同步公文目錄到桌面審核表")
    ap.add_argument("--path", help="審核表路徑(預設讀 env.env review_sheet_path,再預設桌面)")
    ap.add_argument("--overwrite", action="store_true", help="連非空白欄一起覆寫(預設不覆寫)")
    ap.add_argument("--list", metavar="GATE", choices=GATES, help="列出該關卡已打 OK 的列")
    ap.add_argument("--done", nargs="+", metavar="文號",
                    help="把這些公文的兩道關卡標成「已辦」（你自己辦完的）")
    ap.add_argument("--gate", choices=GATES,
                    help="搭配 --done：只標某一關（預設兩關都標）")
    ap.add_argument("--prepare", action="store_true",
                    help="先對還沒備料的公文補跑 LLM 摘要，再同步（會花 token）")
    a = ap.parse_args()

    p = a.path or sheet_path()
    if a.done:
        gates = [a.gate] if a.gate else list(GATES)
        recs = [{DOC_NO_COL: n.strip(), **{g: DONE_MARK for g in gates}}
                for n in a.done]
        try:
            upsert(recs, path=p, overwrite=True)
        except PermissionError as e:
            print(f"[review_sheet] {e}")
            raise SystemExit(1)
        print(f"[review_sheet] 已把 {len(recs)} 筆的「{'／'.join(gates)}」標成「{DONE_MARK}」")
        return
    if a.list:
        for r in approved(a.list, p):
            print(f"  [{r['_row']}] {r['文號']}  {(r['主旨'] or '')[:40]}")
        return
    if a.prepare:
        ok, failed = prepare()
        print(f"\n[review_sheet] 備料完成 {ok} 個" +
              (f"，失敗 {len(failed)} 個：{failed}" if failed else ""))
    try:
        added, updated, n, skipped = sync(path=p, overwrite=a.overwrite)
    except PermissionError as e:
        print(f"[review_sheet] {e}")
        raise SystemExit(1)
    print(f"[review_sheet] 掃到 {n} 個公文目錄 → 新增 {added} 列、更新 {updated} 列"
          f"（{skipped} 個尚未備料，略過）")
    print(f"[review_sheet] {p}")

    # 「什麼都沒發生」必須看得出來是正常還是有事要做 ——
    # 2026-07-28 實測:印「新增 0 列」被當成程式壞掉,使用者改跑 main.py，
    # 結果結案流程自動把一份內部公文貼上校網。零結果一定要附下一步。
    if added == 0 and updated == 0:
        print()
        if skipped:
            print(f"  ⓘ 沒有新資料進來，因為那 {skipped} 個目錄還沒跑摘要。")
            print(f"     要補摘要請跑（會叫 LLM、花 token）：")
            print(f"         python review_sheet.py --prepare")
        else:
            print("  ⓘ 表已是最新，沒有需要更新的欄位。這是正常的，不是出錯。")
        print("     還沒下載新公文的話，先跑 py main.py 4（備料）。")
        print("     ⚠️ 不要跑 py main.py 或 py main.py 3 —— 那兩條會自動送簽／自動貼校網。")


if __name__ == "__main__":
    main()
